# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2014-2026 Mohsen Hozan
#
# This file is part of Emognition. It is free software: you can redistribute
# it and/or modify it under the terms of the GNU Affero General Public License
# as published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version. See the LICENSE file.

"""The four report formats must agree with each other (and with Excel's rounding)."""

import csv

import pytest
from openpyxl import load_workbook

from emognition.core import Emotion, Response, Sex, Stimulus, summarize
from emognition.reports import SessionInfo, write_csv_report, write_excel_report
from emognition.reports.common import (
    EMPTY,
    fmt_pct,
    fmt_pct_unit,
    fmt_replays,
    fmt_seconds,
    fmt_seconds_unit,
    prepare_summary,
)


@pytest.mark.parametrize("fmt, value, text", [
    (fmt_pct, 6.25, "6.3"),        # Python's f"{6.25:.1f}" gives 6.2; Excel shows 6.3
    (fmt_pct, 62.5, "62.5"),
    (fmt_seconds, 0.0005, "0.001"),
    (fmt_seconds, 2.675, "2.675"),
    (fmt_replays, 0.125, "0.13"),  # f"{0.125:.2f}" gives 0.12
])
def test_ties_round_half_up_like_excel(fmt, value, text):
    assert fmt(value) == text


def test_missing_values_never_get_a_unit():
    assert fmt_pct_unit(None) == EMPTY and fmt_seconds_unit(None) == EMPTY
    assert fmt_pct_unit(50.0) == "50.0%" and fmt_seconds_unit(0.5) == "0.500 s"


def _responses():
    stim = lambda emo, n: Stimulus(path=f"/c/{n}.wav", emotion=emo, sex=Sex.MALE)  # noqa: E731
    return [
        Response(stimulus=stim(Emotion.SAD, i), responded=Emotion.SAD,
                 rt_onset=1.5 + 0.01 * i, rt_offset=0.5 + 0.01 * i)
        for i in range(8)
    ] + [Response(stimulus=stim(Emotion.SAD, 9), responded=Emotion.SAD, rt_onset=40.0, rt_offset=39.0)]


def test_summary_csv_counts_the_rts_behind_each_mean(tmp_path):
    _, summary_path = write_csv_report(_responses(), str(tmp_path / "r.csv"))
    rows = {row["group"]: row for row in csv.DictReader(open(summary_path, newline="", encoding="utf-8"))}
    assert rows["Total"]["n_answered"] == "9" and rows["Total"]["n_rt"] == "8"   # 1 slow outlier left out


def test_excel_says_when_means_exclude_flagged_trials(tmp_path):
    path = write_excel_report(_responses(), str(tmp_path / "r.xlsx"), info=SessionInfo())
    ws = load_workbook(path)["Percentages"]
    notes = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("exclude 1 trial" in n for n in notes)


def test_a_summary_with_the_other_rt_basis_is_rejected():
    onset_summary = summarize(_responses(), rt_kind="onset")
    with pytest.raises(ValueError, match="onset"):
        prepare_summary(_responses(), summary=onset_summary, rt_kind="offset")


# --- cross-format agreement and edge-case sessions (Phase 4) ------------------

def _mixed_session():
    from emognition.core import Emotion as E

    def r(i, true, chosen, rt):
        stim = Stimulus(path=f"/c/{i}_{true.tag}_f.wav", emotion=true, sex=Sex.FEMALE)
        if chosen is None:
            return Response(stimulus=stim)
        return Response(stimulus=stim, responded=chosen, rt_onset=rt + 1.6, rt_offset=rt)

    plan = [(E.HAPPY, E.HAPPY, 0.41), (E.HAPPY, E.HAPPY, 0.52), (E.HAPPY, E.SAD, 0.93),
            (E.SAD, E.SAD, 0.61), (E.SAD, E.NEUTRAL, 0.77), (E.SAD, E.SAD, 0.58),
            (E.ANGRY, E.ANGRY, 0.66), (E.ANGRY, E.SCARED, 1.05), (E.NEUTRAL, None, None)]
    return [r(i, t, c, rt) for i, (t, c, rt) in enumerate(plan)]


def test_csv_excel_and_html_report_the_same_numbers(tmp_path):
    import re

    from emognition.core import EMOTIONS_IN_ORDER
    from emognition.reports import render_html_report
    from emognition.reports.common import round_half_up

    responses = _mixed_session()
    reference, _flags = prepare_summary(responses, summary=None, rt_kind="offset")
    _, summary_csv = write_csv_report(responses, str(tmp_path / "r.csv"))
    csv_rows = {row["group"]: row for row in csv.DictReader(open(summary_csv, newline="", encoding="utf-8"))}
    ws = load_workbook(write_excel_report(responses, str(tmp_path / "r.xlsx")))["Percentages"]
    excel_rows = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value)
                  for r in range(2, ws.max_row + 1)}
    html = render_html_report(responses)
    stats_html = html.split("<h2>Statistics</h2>", 1)[1].split("</section>", 1)[0]

    for emo in EMOTIONS_IN_ORDER:
        ref = reference.per_emotion[emo]
        pct_x, rt_x = excel_rows[emo.label]
        m = re.search(rf"<th class='rowhead'>{emo.label}</th><td>([^<]*)</td><td>([^<]*)</td>"
                      rf"<td>([^<]*)</td>", stats_html)
        assert m, emo.label
        if ref.pct_correct is None:                          # never answered
            assert csv_rows[emo.label]["pct_correct"] == "" and pct_x is None and m.group(1) == EMPTY
            continue
        assert float(csv_rows[emo.label]["pct_correct"]) == float(round_half_up(ref.pct_correct, 2))
        assert pct_x == pytest.approx(ref.pct_correct)
        assert m.group(1) == fmt_pct(ref.pct_correct)
        assert float(csv_rows[emo.label]["mean_rt_s"]) == float(round_half_up(ref.mean_rt, 4))
        assert rt_x == pytest.approx(ref.mean_rt) and m.group(2) == fmt_seconds(ref.mean_rt)


def _session(kind):
    one = Stimulus(path="/c/01_hap_f.wav", emotion=Emotion.HAPPY, sex=Sex.FEMALE)
    if kind == "empty":
        return []
    if kind == "single":
        return [Response(stimulus=one, responded=Emotion.HAPPY, rt_onset=1.5, rt_offset=0.5)]
    if kind == "all wrong":
        return [Response(stimulus=one, responded=Emotion.SAD, rt_onset=1.5 + i / 10, rt_offset=0.5 + i / 10)
                for i in range(4)]
    return [Response(stimulus=one) for _ in range(3)]          # all unanswered


@pytest.mark.parametrize("kind", ["empty", "single", "all wrong", "all unanswered"])
def test_every_writer_copes_with_edge_case_sessions(tmp_path, kind):
    from emognition.reports import write_html_report, write_pdf_report

    responses = _session(kind)
    base = str(tmp_path / "r")
    trials, summary_csv = write_csv_report(responses, base + ".csv")
    write_excel_report(responses, base + ".xlsx", info=SessionInfo())
    write_html_report(responses, base + ".html", info=SessionInfo())
    write_pdf_report(responses, base + ".pdf", info=SessionInfo())
    total = {row["group"]: row for row in csv.DictReader(open(summary_csv, newline="", encoding="utf-8"))}["Total"]
    expected_answered = {"empty": "0", "single": "1", "all wrong": "4", "all unanswered": "0"}[kind]
    assert total["n_answered"] == expected_answered
    if kind == "all wrong":
        assert float(total["pct_correct"]) == 0.0
    if kind in ("empty", "all unanswered"):
        assert total["pct_correct"] == "" and total["mean_rt_s"] == ""


# --- new measures in the reports (Phase 5) ------------------------------------------

def test_reports_carry_hu_kappa_spread_and_background(tmp_path):
    from emognition.reports import render_html_report, write_confusion_csv

    responses = _mixed_session()
    responses[0].background = "Starry night"
    trials, summary_csv = write_csv_report(responses, str(tmp_path / "r.csv"))
    rows = {row["group"]: row for row in csv.DictReader(open(summary_csv, newline="", encoding="utf-8"))}
    assert rows["Happy"]["hu"] and rows["Total"]["kappa"] and rows["Total"]["median_rt_s"]
    assert rows["Female"]["hu"] == "" and rows["Happy"]["kappa"] == ""
    first = next(csv.DictReader(open(trials, newline="", encoding="utf-8")))
    assert first["background"] == "Starry night"

    long = list(csv.DictReader(open(write_confusion_csv(responses, str(tmp_path / "c.csv")),
                                    newline="", encoding="utf-8")))
    assert len(long) == 25 and {"true_emotion", "chosen_emotion", "count", "pct_of_row"} <= set(long[0])
    happy = {r["chosen_emotion"]: r for r in long if r["true_emotion"] == "Happy"}
    assert happy["Happy"]["count"] == "2" and float(happy["Happy"]["pct_of_row"]) == pytest.approx(66.67)

    wb = load_workbook(write_excel_report(responses, str(tmp_path / "r.xlsx")))
    info = {wb["Session Info"].cell(row=r, column=1).value: wb["Session Info"].cell(row=r, column=2).value
            for r in range(1, wb["Session Info"].max_row + 1)}
    assert info["Cohen's kappa"] is not None
    assert wb["Percentages"].cell(row=1, column=8).value == "Unbiased hit rate (Hu)"
    assert wb["Confusion (% of row)"].cell(row=2, column=2).value == pytest.approx(200 / 3)

    from html import unescape

    html = unescape(render_html_report(responses))     # labels are HTML-escaped
    assert "Cohen's κ" in html and "Unbiased hit rate (Hu)" in html and "66.7%" in html
