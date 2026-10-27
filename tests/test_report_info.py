# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2014-2026 Mohsen Hozan
#
# This file is part of Emognition. It is free software: you can redistribute
# it and/or modify it under the terms of the GNU Affero General Public License
# as published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version. See the LICENSE file.

"""The session-details block is shared by every report format (reports/common.py)."""

from datetime import datetime

from openpyxl import load_workbook

from emognition.core import Emotion, Response, Sex, Stimulus
from emognition.reports import SessionInfo, render_html_report, write_excel_report
from emognition.reports.common import EMPTY, fmt_duration, info_rows

_INFO = SessionInfo(
    participant="P01", experimenter="Tester", mode="Test", audio_folder="/clips",
    background="Slideshow: 5 pictures, every 5 trials, shuffled",
    started_at=datetime(2026, 10, 27, 14, 0, 0), finished_at=datetime(2026, 10, 27, 15, 2, 5),
    app_version="3.0.0", notes="pilot", extra={"Seed": "42"},
)


def _responses():
    stim = Stimulus(path="/clips/01_hap_f.wav", emotion=Emotion.HAPPY, sex=Sex.FEMALE)
    return [Response(stimulus=stim, responded=Emotion.HAPPY, rt_onset=1.6, rt_offset=0.6)]


def test_fmt_duration_uses_matlab_style_hh_mm_ss():
    assert fmt_duration(_INFO.started_at, _INFO.finished_at) == "01:02:05"
    assert fmt_duration(None, _INFO.finished_at) == EMPTY


def test_info_rows_order_and_content():
    rows = dict(info_rows(_INFO, "offset"))
    assert list(rows)[:5] == ["Participant", "Experimenter", "Mode", "Audio folder", "Background"]
    assert rows["Background"].startswith("Slideshow")
    assert rows["Duration"] == "01:02:05" and rows["Reaction-time basis"] == "offset"
    assert rows["Seed"] == "42" and rows["Notes"] == "pilot"


def test_html_and_excel_show_the_same_details(tmp_path):
    html = render_html_report(_responses(), info=_INFO)
    path = write_excel_report(_responses(), str(tmp_path / "r.xlsx"), info=_INFO)
    ws = load_workbook(path)["Session Info"]
    excel = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
             for r in range(1, ws.max_row + 1)}
    for label, value in info_rows(_INFO, "offset"):
        assert label in html and value in html
        assert excel[label] == value
