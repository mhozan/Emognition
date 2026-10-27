# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2014-2026 Mohsen Hozan
#
# This file is part of Emognition. It is free software: you can redistribute
# it and/or modify it under the terms of the GNU Affero General Public License
# as published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version. See the LICENSE file.

"""Exported text must never run as a spreadsheet formula (CSV/Excel injection)."""

import csv

import pytest
from openpyxl import load_workbook

from emognition.core import Emotion, Response, Sex, Stimulus
from emognition.io.spreadsheet import neutralize, restore
from emognition.naming import manifest
from emognition.reports import SessionInfo, write_csv_report, write_excel_report

_EVIL = '=HYPERLINK("http://example.invalid","click")'


@pytest.mark.parametrize("text", ["=1+1", "+1", "-2", "@SUM(A1)", "\tx", "\rx"])
def test_neutralize_prefixes_formula_triggers_and_restore_undoes_it(text):
    assert neutralize(text) == "'" + text
    assert restore(neutralize(text)) == text


@pytest.mark.parametrize("value", ["P01", "happy.wav", "", 3.5, None, "it's fine"])
def test_neutralize_leaves_harmless_values_alone(value):
    assert neutralize(value) == value
    if isinstance(value, str):
        assert restore(value) == value


def _responses(filename):
    stim = Stimulus(path=f"/clips/{filename}", emotion=Emotion.HAPPY, sex=Sex.FEMALE,
                    recording_id=filename.split("_")[0])
    return [Response(stimulus=stim, responded=Emotion.SAD, rt_onset=1.5, rt_offset=0.5)]


def test_excel_stores_formula_like_text_as_text(tmp_path):
    path = write_excel_report(_responses(f"{_EVIL}_hap_f.wav"), str(tmp_path / "r.xlsx"),
                              info=SessionInfo(participant=_EVIL, notes="=cmd|' /C calc'!A0"))
    wb = load_workbook(path)
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert all(c.data_type != "f" for c in cells)                 # no live formulas anywhere
    assert any(c.value == _EVIL for c in cells)                   # text preserved exactly


def test_csv_prefixes_formula_like_file_names(tmp_path):
    trials, _summary = write_csv_report(_responses("=cmd_hap_f.wav"), str(tmp_path / "r.csv"))
    with open(trials, newline="", encoding="utf-8") as fh:
        row = next(csv.DictReader(fh))
    assert row["audio_file"] == "'=cmd_hap_f.wav" and row["recording_id"] == "'=cmd"


def test_manifest_round_trips_formula_like_names(tmp_path):
    (tmp_path / "=bad_hap_f.wav").write_bytes(b"")
    records = manifest.build_manifest(str(tmp_path))
    path = manifest.write_manifest(records, str(tmp_path / "stimuli.tsv"))
    raw = open(path, encoding="utf-8").read()
    assert "'=bad_hap_f.wav" in raw                               # inert in a spreadsheet
    assert [r.filename for r in manifest.read_manifest_tsv(path)] == ["=bad_hap_f.wav"]
