# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2014-2026 Mohsen Hozan
#
# This file is part of Emognition. It is free software: you can redistribute
# it and/or modify it under the terms of the GNU Affero General Public License
# as published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version. See the LICENSE file.

"""Excel (.xlsx) report writer — the FOSS replacement for MATLAB's writematrix.

Sheet layout mirrors the original ``savestate.m`` output so existing analysis
pipelines keep working, plus a modern "Session Info" summary sheet:

    1. Session Info     metadata + headline accuracy
    2. Confusion Matrix rows = true emotion, cols = chosen (heat-mapped)
    3. Percentages      per-emotion / female / male / Total x (%correct, RT, replays)
    4. Pressed Keys     one row per trial + a 1->Happy..5->Angry legend

openpyxl crash course (the bits used here):
  * ``Workbook()`` starts with one sheet (``wb.active``); ``wb.create_sheet(name)``
    adds more. Cells are **1-indexed**: ``ws.cell(row=1, column=1)`` is "A1".
  * Writing a Python ``float``/``int``/``bool`` keeps it a real number/boolean in
    Excel (so formulas work); ``None`` leaves the cell blank. We write raw numbers
    and attach a display ``number_format`` rather than pre-formatted strings.
  * Styling is per-cell: ``cell.font``, ``cell.alignment``, ``cell.fill``.
"""

from __future__ import annotations

from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..core import EMOTIONS_IN_ORDER, SEXES_IN_ORDER, Response, RtFlag, SessionSummary
from ..core.scoring import DEFAULT_RT_KIND, GroupStats
from .common import (
    RT_FLAG_LABEL,
    STAT_COLUMNS,
    SessionInfo,
    heat_rgb,
    info_rows,
    n_rt_excluded,
    prepare_summary,
)

_BOLD = Font(bold=True)
_CENTER = Alignment(horizontal="center")
_TITLE = Font(bold=True, size=14)


# --------------------------------------------------------------------------
# small styling helpers
# --------------------------------------------------------------------------

def _text_cell(ws: Worksheet, row: int, column: int, value):
    """Write *value* so text is always stored as text, never evaluated as a formula.

    openpyxl turns any string starting with ``=`` into a live formula; a
    participant label or stimulus file name must not be able to do that
    (spreadsheet formula injection).
    """
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and cell.data_type == "f":
        cell.data_type = "s"
    return cell


def _header(cell) -> None:
    """Make a cell a bold, centred header."""
    cell.font = _BOLD
    cell.alignment = _CENTER


def _cell_fill(count: int, row_total: int, on_diagonal: bool) -> PatternFill | None:
    """Heat-map fill for a confusion cell (shared shading: share of its row)."""
    rgb = heat_rgb(count, row_total, on_diagonal)
    if rgb is None:
        return None
    hexcol = "{:02X}{:02X}{:02X}".format(*rgb)
    return PatternFill(start_color=hexcol, end_color=hexcol, fill_type="solid")


# --------------------------------------------------------------------------
# individual sheets
# --------------------------------------------------------------------------

def _write_session_info(
    ws: Worksheet, info: SessionInfo, summary: SessionSummary, n_trials: int
) -> None:
    """Metadata + headline numbers (the modern addition to the legacy report)."""
    ws["A1"] = "Emognition — Session Report"
    ws["A1"].font = _TITLE

    rt_label = f"reaction time basis: {summary.rt_kind}"
    rows: list[tuple[str, object]] = [
        *info_rows(info, summary.rt_kind),
        ("Trials presented", summary.total.n_presented),
        ("Trials answered", summary.total.n_answered),
        ("Overall correct", summary.total.n_correct),
        ("Overall accuracy (%)", summary.total.pct_correct),
        ("Mean reaction (s)", summary.total.mean_rt),
    ]
    number_formats = {"Overall accuracy (%)": "0.0", "Mean reaction (s)": "0.000"}
    for r, (label, value) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=label).font = _BOLD
        cell = _text_cell(ws, r, 2, value)
        if label in number_formats:
            cell.number_format = number_formats[label]
    ws.cell(row=len(rows) + 4, column=1, value=rt_label).font = Font(italic=True)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 32


def _write_confusion(ws: Worksheet, summary: SessionSummary) -> None:
    """5x5 confusion matrix, rows = true emotion, cols = chosen emotion."""
    labels = [e.label for e in summary.emotion_order]
    corner = ws.cell(row=1, column=1, value="true \\ chosen")
    _header(corner)

    for j, lab in enumerate(labels, start=2):      # column headers (chosen)
        _header(ws.cell(row=1, column=j, value=lab))
    for i, lab in enumerate(labels, start=2):      # row headers (true)
        _header(ws.cell(row=i, column=1, value=lab))

    for i, row in enumerate(summary.confusion):
        for j, count in enumerate(row):
            cell = ws.cell(row=i + 2, column=j + 2, value=count)
            cell.alignment = _CENTER
            fill = _cell_fill(count, sum(row), on_diagonal=(i == j))
            if fill is not None:
                cell.fill = fill

    ws.column_dimensions["A"].width = 14
    ws.freeze_panes = "B2"  # keep labels visible when scrolling


def _stat_row(ws: Worksheet, row: int, label: str, stats: GroupStats) -> None:
    """Write one (% correct, mean RT, mean replays) row of the Percentages sheet."""
    ws.cell(row=row, column=1, value=label).font = _BOLD
    pct = ws.cell(row=row, column=2, value=stats.pct_correct)
    rt = ws.cell(row=row, column=3, value=stats.mean_rt)
    rep = ws.cell(row=row, column=4, value=stats.mean_replays)
    pct.number_format = "0.0"
    rt.number_format = "0.000"
    rep.number_format = "0.00"


def _write_percentages(ws: Worksheet, summary: SessionSummary, n_excluded: int = 0) -> None:
    """Per-emotion, per-sex, and total stats (legacy 'Percentages' sheet)."""
    for j, title in enumerate(STAT_COLUMNS, start=2):
        _header(ws.cell(row=1, column=j, value=title))

    row = 2
    for emo in summary.emotion_order:
        _stat_row(ws, row, emo.label, summary.per_emotion[emo])
        row += 1
    for sex in SEXES_IN_ORDER:
        _stat_row(ws, row, sex.label, summary.per_sex[sex]); row += 1
    _stat_row(ws, row, "Total", summary.total)
    if n_excluded:
        # Say so where the means are shown (HTML and PDF already did).
        note = (f"Mean reaction times exclude {n_excluded} trial(s) flagged as "
                "anticipatory or slow outliers (see 'RT flag' in Pressed Keys).")
        ws.cell(row=row + 2, column=1, value=note).font = Font(italic=True)

    ws.column_dimensions["A"].width = 12
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 17
    ws.freeze_panes = "A2"


def _write_pressed_keys(
    ws: Worksheet,
    responses: Sequence[Response],
    flags: Sequence[RtFlag],
    rt_kind: str,
) -> None:
    """One row per trial, plus a code legend (legacy 'Pressed Keys' sheet)."""
    headers = [
        "Trial", "Audio filename", "Played (true)", "Pressed (chosen)",
        "Correct?", "RT onset (s)", "RT offset (s)", "Replays", "RT flag",
    ]
    for j, h in enumerate(headers, start=1):
        _header(ws.cell(row=1, column=j, value=h))

    for i, (r, flag) in enumerate(zip(responses, flags), start=1):
        ws.cell(row=i + 1, column=1, value=i)
        _text_cell(ws, i + 1, 2, r.stimulus.filename)
        ws.cell(row=i + 1, column=3, value=int(r.true_emotion))
        ws.cell(row=i + 1, column=4, value=None if r.responded is None else int(r.responded))
        ws.cell(row=i + 1, column=5, value=r.is_correct if r.answered else None)
        c6 = ws.cell(row=i + 1, column=6, value=r.rt_onset)
        c7 = ws.cell(row=i + 1, column=7, value=r.rt_offset)
        c6.number_format = c7.number_format = "0.000"
        ws.cell(row=i + 1, column=8, value=r.n_replays)
        ws.cell(row=i + 1, column=9, value=RT_FLAG_LABEL[flag] or None)

    # Legend (column K), explaining the integer codes used above.
    _header(ws.cell(row=1, column=11, value="Legend"))
    for k, emo in enumerate(EMOTIONS_IN_ORDER, start=2):
        ws.cell(row=k, column=11, value=f"{int(emo)}: {emo.label}")

    ws.column_dimensions["B"].width = 18
    for col in ("C", "D", "E", "F", "G", "I"):
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["K"].width = 14
    ws.freeze_panes = "A2"


# --------------------------------------------------------------------------
# public entry point
# --------------------------------------------------------------------------

def write_excel_report(
    responses: Sequence[Response],
    path: str,
    *,
    summary: SessionSummary | None = None,
    info: SessionInfo | None = None,
    rt_kind: str = DEFAULT_RT_KIND,
) -> str:
    """Write the full multi-sheet Excel report and return the path.

    Args:
        responses: The session's completed trials (one row each in Pressed Keys).
        path: Output ``.xlsx`` path.
        summary: Precomputed summary; if None, computed from ``responses``.
        info: Session metadata for the header sheet; defaults to blanks.
        rt_kind: RT basis for the summary if it must be computed ("offset" default).
    """
    summary, flags = prepare_summary(responses, summary=summary, rt_kind=rt_kind)
    info = info or SessionInfo(rt_kind=rt_kind)

    wb = Workbook()
    info_ws = wb.active
    info_ws.title = "Session Info"
    _write_session_info(info_ws, info, summary, len(responses))
    _write_confusion(wb.create_sheet("Confusion Matrix"), summary)
    _write_percentages(wb.create_sheet("Percentages"), summary, n_rt_excluded(flags))
    _write_pressed_keys(wb.create_sheet("Pressed Keys"), responses, flags, rt_kind)

    wb.save(path)
    return path
